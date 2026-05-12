#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
镍电解工智能备考系统
====================
运行方式:  streamlit run app.py
依赖安装:  pip install streamlit python-docx

功能：
  - 解析 .docx 题库，自动推断题型与类别
  - 残缺题拦截（题干过短 / 无答案 / 答案无效）
  - 智能判分（填空模糊匹配 / 选择防乱序 / 判断多格式）
  - 首页大盘 / 智能刷题 / 错题本
"""

import io
import re
import difflib
import random
from collections import defaultdict
from typing import Optional, Tuple, List, Dict, Any

import streamlit as st

# ── 检查可选依赖 ──────────────────────────────────────────────────────────────
try:
    from docx import Document
    DOCX_OK = True
except ImportError:
    DOCX_OK = False


# ╔══════════════════════════════════════════════════════════════════════════════╗
# ║  第一模块：文档解析与残缺题拦截                                             ║
# ╚══════════════════════════════════════════════════════════════════════════════╝

# 题号行正则：匹配 "1、" "2." "3．" "4。" "5 " 等各种分隔符
_Q_START = re.compile(r'^(\d{1,4})\s*[、.．。\s]\s*(.*)', re.DOTALL)

# 答案行正则：匹配 "答：" "答案：" "答案:" 等
_ANS_LINE = re.compile(r'^答案?\s*[：:]\s*(.*)', re.DOTALL)


def _iter_all_paragraphs(doc_bytes: bytes) -> List[str]:
    """
    从 docx 字节流中按文档顺序提取所有非空文本行。
    优先遍历正文段落；若段落中嵌有表格则一并提取（有些题库用表格排版）。
    """
    doc = Document(io.BytesIO(doc_bytes))
    lines: List[str] = []

    # 通过 XML body 子元素保持段落与表格的原始顺序
    from docx.oxml.ns import qn
    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag

        if tag == 'p':                            # 正文段落
            text = child.text_content().strip() if hasattr(child, 'text_content') else ''
            # python-docx 提供更可靠的 .text
            try:
                from docx.oxml import CT_P
                from docx.text.paragraph import Paragraph
                para = Paragraph(child, doc)
                text = para.text.strip()
            except Exception:
                pass
            if text:
                lines.append(text)

        elif tag == 'tbl':                        # 表格
            from docx.table import Table
            try:
                tbl = Table(child, doc)
                for row in tbl.rows:
                    for cell in row.cells:
                        for para in cell.paragraphs:
                            t = para.text.strip()
                            if t:
                                lines.append(t)
            except Exception:
                pass

    # 回退：如果上面提取为空（XML API 差异），直接用 doc.paragraphs
    if not lines:
        lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    return lines


def _check_invalid(num: int, text: str, answer: Optional[str]) -> Optional[str]:
    """
    残缺题检查。
    返回原因字符串表示需要拦截；返回 None 表示题目有效。
    """
    # 题干：去掉所有非汉字/字母/数字后，净长不足3字符视为空题干
    clean_text = re.sub(r'[^\w一-鿿]', '', text or '')
    if len(clean_text) < 3:
        return f'题干过短（净长 {len(clean_text)} 字）'

    # 答案缺失
    if answer is None:
        return '缺少答案'

    stripped = answer.strip()
    if not stripped:
        return '答案为空'

    # 答案为"略/无/见课本"等无效占位词
    if re.fullmatch(
        r'(略|无|暂略|见课本|待定|参考答案|[\(（\)）\s/\\]+)',
        stripped
    ):
        return f'答案无效（值={stripped!r}）'

    return None


def infer_type(text: str, answer: str) -> str:
    """根据题干结构和答案内容推断题型。"""
    ans = answer.strip()

    # ① 判断题：答案是布尔类词
    if re.fullmatch(
        r'(对|错|√|×|T|F|true|false|正确|错误|是|否)',
        ans, re.IGNORECASE
    ):
        return '判断题'

    # ② 选择题/多选题：题干含 A/B/C/D 选项 或 答案是纯字母组合
    has_opts = bool(
        re.search(r'(?m)(?:^|\s)[A-D]\s*[、.．\s。]', text) or
        re.search(r'[（(][A-D][)）]', text)
    )
    ans_letters = re.fullmatch(r'[A-Da-d]{1,4}', ans)

    if has_opts or ans_letters:
        # 题干含"多选"或答案多于1个字母 → 多选题
        if '多选' in text or len(re.findall(r'[A-Da-d]', ans)) > 1:
            return '多选题'
        return '选择题'

    return '填空题'


def infer_category(text: str) -> str:
    """根据关键词为题目打上类别标签（四分类，兜底为基础理论）。"""
    if re.search(
        r'安全|危险|着火|急救|液碱|酸雾|侧上风|高处作业|灭火|报警|漏电|盲板',
        text
    ):
        return '安全与环保类'

    if re.search(
        r'疙瘩|弯板|气孔|钝化|烧板|短路|下饺子|阳极泥|打火|平板作业|掏槽',
        text
    ):
        return '现场异常排查类'

    if re.search(
        r'\d+%|\d+℃|\d+mm|\d+V|\d+A|\d+天|密度|单耗|大于|小于',
        text
    ):
        return '工艺参数硬指标'

    return '基础理论与常识'


def parse_docx(doc_bytes: bytes) -> Tuple[List[Dict], List[Dict]]:
    """
    主解析函数：从 docx 中提取题目，清洗并打标。

    Returns
    -------
    questions : 有效题目列表，每项包含 num/text/answer/type/category
    error_log : 被拦截的残缺题列表，每项包含 num/text/reason
    """
    lines = _iter_all_paragraphs(doc_bytes)

    # ── 状态机：逐行扫描，积累题目块 ────────────────────────────
    raw_blocks: List[Dict] = []

    cur_num: Optional[int] = None
    cur_lines: List[str] = []
    cur_answer: Optional[str] = None
    in_answer = False

    for line in lines:
        q_m = _Q_START.match(line)
        a_m = _ANS_LINE.match(line)

        if q_m:
            # 遇到新题号 → 先保存上一题
            if cur_num is not None:
                raw_blocks.append({
                    'num':    cur_num,
                    'text':   '\n'.join(cur_lines).strip(),
                    'answer': cur_answer,
                })
            cur_num    = int(q_m.group(1))
            first_line = q_m.group(2).strip()
            cur_lines  = [first_line] if first_line else []
            cur_answer = None
            in_answer  = False

        elif a_m and cur_num is not None:
            # 遇到答案标记行
            in_answer  = True
            cur_answer = a_m.group(1).strip()

        elif cur_num is not None:
            if in_answer:
                # 答案续行
                cur_answer = (cur_answer or '') + '\n' + line
            else:
                # 题干续行
                cur_lines.append(line)

    # 别忘最后一道题
    if cur_num is not None:
        raw_blocks.append({
            'num':    cur_num,
            'text':   '\n'.join(cur_lines).strip(),
            'answer': cur_answer,
        })

    # ── 过滤 + 打标 ──────────────────────────────────────────────
    questions: List[Dict] = []
    error_log: List[Dict] = []

    for b in raw_blocks:
        reason = _check_invalid(b['num'], b['text'], b['answer'])
        if reason:
            error_log.append({
                'num':    b['num'],
                'text':   (b['text'] or '')[:60],
                'reason': reason,
            })
            continue

        ans_clean = (b['answer'] or '').strip()
        q_type    = infer_type(b['text'], ans_clean)
        category  = infer_category(b['text'])

        questions.append({
            'num':      b['num'],
            'text':     b['text'],
            'answer':   ans_clean,
            'type':     q_type,
            'category': category,
        })

    return questions, error_log


# ╔══════════════════════════════════════════════════════════════════════════════╗
# ║  第二模块：高容错智能判分引擎                                               ║
# ╚══════════════════════════════════════════════════════════════════════════════╝

# 判断题布尔映射（支持多种输入格式）
_BOOL_MAP: Dict[str, bool] = {
    '对': True,  '√': True,  'T': True,  't': True,
    'true': True,  'TRUE': True,  '正确': True,  '是': True,
    '错': False, '×': False, 'F': False, 'f': False,
    'false': False, 'FALSE': False, '错误': False, '否': False,
}

# 用于填空题清洗的标点集合（全角 + 半角）
_PUNCT_RE = re.compile(
    r'[\s　，、．。！？：；'
    r'《》【】‘’“”'
    r',.:;!?\'"()\[\]{}、。，；：！？《》【】]'
)


def _clean(s: str) -> str:
    """去掉所有空白与标点，用于填空题比对。"""
    return _PUNCT_RE.sub('', s)


def _judge_truefalse(user: str, correct: str) -> Tuple[bool, str]:
    """判断题：统一映射为 bool 后比较。"""
    u = _BOOL_MAP.get(user.strip())
    c = _BOOL_MAP.get(correct.strip())

    if u is None:
        return False, f'⚠️ 无法识别输入 "{user}"，请输入：对/错/√/×/T/F'
    if c is None:
        # 标准答案格式异常，尽可能宽容
        return False, f'⚠️ 标准答案格式异常：{correct}'

    return (True, '✅ 回答正确！') if (u == c) else (False, '❌ 回答错误')


def _judge_choice(user: str, correct: str) -> Tuple[bool, str]:
    """
    选择题：提取字母 → 大写 → 去重 → 排序后比对。
    防止用户输入 "BCA" 而标准答案是 "ABC" 被误判。
    """
    u_letters = ''.join(sorted(set(re.findall(r'[A-Da-d]', user)))).upper()
    c_letters = ''.join(sorted(set(re.findall(r'[A-Da-d]', correct)))).upper()

    if not u_letters:
        return False, f'⚠️ 无法识别选项 "{user}"，请输入字母（如 A 或 ABC）'

    return (True, '✅ 回答正确！') if (u_letters == c_letters) else (False, '❌ 回答错误')


def _judge_fillblank(user: str, correct: str) -> Tuple[bool, str]:
    """
    填空题三级容错：
      1. 精确匹配（去标点空格后）
      2. 忽略大小写
      3. difflib 模糊匹配，相似度 >= 0.75 视为正确
    """
    u = _clean(user)
    c = _clean(correct)

    if u == c:
        return True, '✅ 回答正确！'

    if u.lower() == c.lower():
        return True, '✅ 回答正确！（忽略大小写）'

    ratio = difflib.SequenceMatcher(None, u, c).ratio()

    if ratio >= 0.75:
        return True, f'✅ 模糊匹配通过（相似度 {ratio:.0%}）'

    if ratio >= 0.50:
        return False, f'❌ 接近但不够准确（相似度 {ratio:.0%}），再想想'

    return False, '❌ 回答错误'


def smart_judge(
    user_ans: str, correct_ans: str, q_type: str
) -> Tuple[bool, str]:
    """
    智能判分统一入口。
    返回 (是否正确, 反馈文字)
    """
    user_ans = (user_ans or '').strip()
    if not user_ans:
        return False, '⚠️ 未作答'

    if q_type == '判断题':
        return _judge_truefalse(user_ans, correct_ans)
    if q_type in ('选择题', '多选题'):
        return _judge_choice(user_ans, correct_ans)
    return _judge_fillblank(user_ans, correct_ans)


# ╔══════════════════════════════════════════════════════════════════════════════╗
# ║  第三模块：Streamlit 前端                                                   ║
# ╚══════════════════════════════════════════════════════════════════════════════╝

_PAGE_HOME     = '🏠 首页大盘'
_PAGE_PRACTICE = '📖 智能刷题'
_PAGE_WRONG    = '📚 错题本'

_CAT_ICON: Dict[str, str] = {
    '安全与环保类':   '🔴',
    '现场异常排查类': '🟠',
    '工艺参数硬指标': '🔵',
    '基础理论与常识': '🟢',
}
_TYPE_COLOR: Dict[str, str] = {
    '填空题': '#27ae60',
    '选择题': '#2980b9',
    '多选题': '#e67e22',
    '判断题': '#8e44ad',
}


# ── Session State 初始化 ──────────────────────────────────────────────────────

def _init_session() -> None:
    defaults: Dict[str, Any] = {
        'questions':   [],
        'error_log':   [],
        'wrong_book':  [],
        'loaded':      False,
        '_file_key':   '',       # 用于检测文件是否变更
        'cur_q':       None,
        'show_ans':    False,
        'submitted':   False,
        'last_result': None,
    }
    for k, v in defaults.items():
        if k not in st.session_state:
            st.session_state[k] = v


# ── 工具函数 ──────────────────────────────────────────────────────────────────

def _add_to_wrong_book(q: Dict, user_ans: str) -> None:
    """加入错题本，同一题号只记录一次。"""
    wb: List[Dict] = st.session_state.wrong_book
    if not any(item['num'] == q['num'] for item in wb):
        st.session_state.wrong_book.append({**q, 'user_answer': user_ans})


def _pick_question(pool: List[Dict]) -> Dict:
    """从题池随机抽取，尽量避免连续出同一题。"""
    if len(pool) == 1:
        return pool[0]
    last_num = (st.session_state.cur_q or {}).get('num', -1)
    candidates = [q for q in pool if q['num'] != last_num] or pool
    return random.choice(candidates)


def _type_badge_html(q_type: str) -> str:
    color = _TYPE_COLOR.get(q_type, '#888')
    return (
        f'<span style="background:{color};color:#fff;padding:2px 9px;'
        f'border-radius:4px;font-size:12px;font-weight:bold">{q_type}</span>'
    )


def _cat_badge_html(category: str) -> str:
    icon = _CAT_ICON.get(category, '⬜')
    return (
        f'<span style="background:#f0f0f0;color:#444;padding:2px 9px;'
        f'border-radius:4px;font-size:12px">{icon} {category}</span>'
    )


def _gen_export_text(wb: List[Dict]) -> bytes:
    """生成错题本 TXT 内容，使用 UTF-8 BOM 确保 Windows Notepad 兼容。"""
    lines = [
        '镍电解工智能备考系统 · 错题本导出',
        f'共 {len(wb)} 道错题',
        '=' * 60,
        '',
    ]
    for i, q in enumerate(wb, 1):
        icon = _CAT_ICON.get(q['category'], '')
        lines += [
            f"第 {i} 题（原题号 {q['num']}）",
            f"题型：{q['type']}  |  类别：{icon}{q['category']}",
            f"题目：{q['text']}",
            f"你的答案：{q.get('user_answer', '')}",
            f"正确答案：{q['answer']}",
            '-' * 60,
            '',
        ]
    return '\n'.join(lines).encode('utf-8-sig')  # BOM for Windows


# ── 侧边栏 ────────────────────────────────────────────────────────────────────

def _render_sidebar() -> str:
    """渲染侧边栏，返回当前选中的页面名称。"""
    with st.sidebar:
        st.markdown('## 🔬 镍电解工备考系统')
        st.divider()

        # 依赖检查提示
        if not DOCX_OK:
            st.error('缺少依赖！\n\n`pip install python-docx`')
            st.stop()

        # 题库上传
        st.markdown('### 📁 加载题库')
        uploaded = st.file_uploader('上传 .docx 题库文档', type=['docx'])

        if uploaded is not None:
            # 用文件名+大小作为变更检测键，避免重复解析
            file_key = f'{uploaded.name}|{uploaded.size}'
            if st.session_state['_file_key'] != file_key:
                with st.spinner('正在解析题库，请稍候…'):
                    try:
                        qs, errs = parse_docx(uploaded.read())
                        st.session_state.questions  = qs
                        st.session_state.error_log  = errs
                        st.session_state.loaded     = True
                        st.session_state.cur_q      = None
                        st.session_state.submitted  = False
                        st.session_state.show_ans   = False
                        st.session_state.last_result = None
                        st.session_state['_file_key'] = file_key
                        st.success(
                            f'✅ 解析完成\n\n'
                            f'有效题目：**{len(qs)}** 道\n\n'
                            f'拦截残缺：**{len(errs)}** 道'
                        )
                    except Exception as exc:
                        st.error(f'解析失败：{exc}')

        if st.session_state.loaded:
            q_count = len(st.session_state.questions)
            st.caption(f'📚 题库已就绪：{q_count} 道有效题目')
            if st.button('🔄 重新上传题库', use_container_width=True):
                for k in ('questions', 'error_log', 'wrong_book', 'cur_q',
                          'submitted', 'show_ans', 'last_result'):
                    st.session_state[k] = [] if k in ('questions', 'error_log', 'wrong_book') else None
                st.session_state.loaded    = False
                st.session_state['_file_key'] = ''
                st.rerun()

        st.divider()

        page = st.radio(
            '页面导航',
            [_PAGE_HOME, _PAGE_PRACTICE, _PAGE_WRONG],
            label_visibility='collapsed',
        )

        st.divider()
        st.caption('v1.0 · 镍电解工备考系统')

    return page


# ── 首页大盘 ──────────────────────────────────────────────────────────────────

def _render_home() -> None:
    st.title('🔬 镍电解工智能备考系统')

    if not st.session_state.loaded:
        st.info('👈 请先在左侧侧边栏上传题库 Word 文档（.docx）')

        with st.expander('📖 系统说明 & 文档格式要求'):
            st.markdown(
                """
**系统功能**

| 功能 | 说明 |
|------|------|
| 自动解析 | 上传 .docx 后自动提取题目/答案，过滤残缺题 |
| 智能分类 | 自动识别题型（填空/选择/多选/判断）与类别标签 |
| 智能刷题 | 按类别/题型筛选，随机出题，智能容错判分 |
| 错题本   | 自动记录错题，支持导出为 TXT |

**智能判分规则**
- **填空题**：剔除标点空格后，`difflib` 模糊匹配，相似度 ≥ 75% 即视为正确
- **选择题**：自动提取并排序字母，防因顺序不同误判
- **判断题**：支持 对/错/√/×/T/F 等多种格式

**文档格式要求**（示例）
```
1、电解槽的工作电压一般控制在多少范围内？
答：3.8～4.2 V

2、下列哪项属于阳极钝化的原因？（单选）
A、电流密度过低  B、电解液温度过高  C、阳极泥堆积  D、以上都不是
答：C

3、发生触电事故时应首先（  ）。对还是错？
答：对
```
                """
            )
        return

    qs:   List[Dict] = st.session_state.questions
    errs: List[Dict] = st.session_state.error_log
    wb:   List[Dict] = st.session_state.wrong_book

    # 顶部指标卡
    c1, c2, c3, c4 = st.columns(4)
    c1.metric('📚 有效题目', len(qs))
    c2.metric('🚫 拦截残缺', len(errs))
    c3.metric('📝 错题本',   len(wb))
    type_set = {q['type'] for q in qs}
    c4.metric('📋 题型种类', len(type_set))

    st.divider()

    col_l, col_r = st.columns(2)

    # 题型分布
    with col_l:
        st.markdown('#### 📊 题型分布')
        type_cnt: Dict[str, int] = defaultdict(int)
        for q in qs:
            type_cnt[q['type']] += 1

        for t, cnt in sorted(type_cnt.items(), key=lambda x: -x[1]):
            pct   = cnt / len(qs) if qs else 0
            color = _TYPE_COLOR.get(t, '#888')
            st.markdown(
                f'<span style="color:{color};font-weight:bold">{t}</span>：{cnt} 题',
                unsafe_allow_html=True,
            )
            st.progress(pct, text=f'{pct:.1%}')

    # 类别分布
    with col_r:
        st.markdown('#### 🏷️ 类别分布')
        cat_cnt: Dict[str, int] = defaultdict(int)
        for q in qs:
            cat_cnt[q['category']] += 1

        for cat, cnt in sorted(cat_cnt.items(), key=lambda x: -x[1]):
            pct  = cnt / len(qs) if qs else 0
            icon = _CAT_ICON.get(cat, '⬜')
            st.markdown(f'{icon} **{cat}**：{cnt} 题')
            st.progress(pct, text=f'{pct:.1%}')

    # 残缺题明细（折叠）
    if errs:
        st.divider()
        with st.expander(f'⚠️ 拦截的残缺题明细（共 {len(errs)} 条）'):
            for e in errs[:60]:
                st.markdown(
                    f'- 第 **{e["num"]}** 题 &nbsp;|&nbsp; '
                    f'原因：`{e["reason"]}` &nbsp;|&nbsp; '
                    f'片段：{e["text"]}',
                    unsafe_allow_html=True,
                )
            if len(errs) > 60:
                st.caption(f'…还有 {len(errs)-60} 条未显示')


# ── 智能刷题 ──────────────────────────────────────────────────────────────────

def _render_input_widget(q: Dict) -> str:
    """根据题型渲染对应输入控件，返回用户原始输入。"""
    key_prefix = f'q{q["num"]}'

    if q['type'] == '判断题':
        val = st.radio(
            '请选择：',
            ['对 ✓', '错 ✗'],
            horizontal=True,
            key=f'{key_prefix}_tf',
        )
        return '对' if '对' in val else '错'

    if q['type'] == '选择题':
        return st.text_input(
            '请输入选项字母（如 A）：',
            key=f'{key_prefix}_sc',
        ).strip()

    if q['type'] == '多选题':
        return st.text_input(
            '请输入选项字母（如 ABC，顺序不限）：',
            key=f'{key_prefix}_mc',
        ).strip()

    # 填空题
    return st.text_area(
        '请输入答案：',
        height=110,
        key=f'{key_prefix}_fb',
    ).strip()


def _render_practice() -> None:
    st.title('📖 智能刷题')

    if not st.session_state.loaded:
        st.warning('👈 请先上传题库文档')
        return

    qs: List[Dict] = st.session_state.questions
    if not qs:
        st.error('题库为空，请检查文档解析结果')
        return

    # ── 筛选器 ──────────────────────────────────────────────────
    all_cats  = ['全部'] + sorted({q['category'] for q in qs})
    all_types = ['全部'] + sorted({q['type']     for q in qs})

    fc1, fc2 = st.columns(2)
    with fc1:
        sel_cat  = st.selectbox('📂 类别筛选', all_cats,  key='sel_cat')
    with fc2:
        sel_type = st.selectbox('📋 题型筛选', all_types, key='sel_type')

    pool = [
        q for q in qs
        if (sel_cat  == '全部' or q['category'] == sel_cat)
        and (sel_type == '全部' or q['type']     == sel_type)
    ]

    st.caption(f'当前筛选范围：**{len(pool)}** 道题目')

    if not pool:
        st.warning('当前筛选条件下没有题目，请调整')
        return

    # ── 抽题按钮 ────────────────────────────────────────────────
    if (
        st.button('🎲 随机抽题', type='primary')
        or st.session_state.cur_q is None
    ):
        st.session_state.cur_q       = _pick_question(pool)
        st.session_state.show_ans    = False
        st.session_state.submitted   = False
        st.session_state.last_result = None
        st.rerun()

    q: Dict = st.session_state.cur_q
    if q is None:
        return

    # ── 题目展示 ────────────────────────────────────────────────
    st.divider()

    # 题型 + 类别徽章
    st.markdown(
        _type_badge_html(q['type']) + '&nbsp;&nbsp;' +
        _cat_badge_html(q['category']) + '&nbsp;&nbsp;' +
        f'<span style="color:#aaa;font-size:12px">第 {q["num"]} 题</span>',
        unsafe_allow_html=True,
    )
    st.markdown('')  # 小间距
    st.markdown(f'### {q["text"]}')

    # ── 答题区（未提交且未查看答案时展示）────────────────────────
    if not st.session_state.submitted and not st.session_state.show_ans:
        user_input = _render_input_widget(q)

        btn_c1, btn_c2 = st.columns(2)
        with btn_c1:
            if st.button('✅ 提交答案', type='primary', use_container_width=True):
                if not user_input:
                    st.warning('请先填写答案')
                else:
                    ok, msg = smart_judge(user_input, q['answer'], q['type'])
                    st.session_state.last_result = {
                        'ok':       ok,
                        'msg':      msg,
                        'user_ans': user_input,
                    }
                    st.session_state.submitted = True
                    if not ok:
                        _add_to_wrong_book(q, user_input)
                    st.rerun()

        with btn_c2:
            if st.button('👁️ 看答案 / 跳过', use_container_width=True):
                st.session_state.show_ans = True
                _add_to_wrong_book(q, '（跳过）')
                st.rerun()

    # ── 判分结果展示 ────────────────────────────────────────────
    if st.session_state.submitted and st.session_state.last_result:
        res = st.session_state.last_result
        if res['ok']:
            st.success(res['msg'])
        else:
            st.error(res['msg'])
            st.info(f'💡 **标准答案：** {q["answer"]}')

        if st.button('➡️ 下一题', type='primary', key='next_after_submit'):
            st.session_state.cur_q       = _pick_question(pool)
            st.session_state.submitted   = False
            st.session_state.show_ans    = False
            st.session_state.last_result = None
            st.rerun()

    # ── 直接看答案后展示 ────────────────────────────────────────
    if st.session_state.show_ans:
        st.info(f'💡 **标准答案：** {q["answer"]}')

        if st.button('➡️ 下一题', type='primary', key='next_after_skip'):
            st.session_state.cur_q       = _pick_question(pool)
            st.session_state.submitted   = False
            st.session_state.show_ans    = False
            st.session_state.last_result = None
            st.rerun()


# ── 错题本 ────────────────────────────────────────────────────────────────────

def _render_wrong_book() -> None:
    st.title('📚 错题本')

    wb: List[Dict] = st.session_state.wrong_book

    if not wb:
        st.info('🎉 错题本为空，继续加油！')
        return

    # 操作栏
    op1, op2, op3 = st.columns([2, 1, 1])
    op1.metric('错题总数', len(wb))

    with op2:
        st.download_button(
            '📥 导出错题本 TXT',
            data=_gen_export_text(wb),
            file_name='错题本.txt',
            mime='text/plain',
            type='primary',
            use_container_width=True,
        )

    with op3:
        if st.button('🗑️ 清空错题本', use_container_width=True):
            st.session_state.wrong_book = []
            st.rerun()

    st.divider()

    # 按类别分组展示
    grouped: Dict[str, List[Dict]] = defaultdict(list)
    for item in wb:
        grouped[item['category']].append(item)

    for cat, items in grouped.items():
        icon = _CAT_ICON.get(cat, '⬜')
        st.markdown(f'#### {icon} {cat}（{len(items)} 题）')

        for item in items:
            preview = item['text'][:45] + ('…' if len(item['text']) > 45 else '')
            with st.expander(
                f'第 {item["num"]} 题 [{item["type"]}]  {preview}'
            ):
                st.markdown(f'**题目：** {item["text"]}')

                ua = item.get('user_answer', '')
                if ua == '（跳过）':
                    st.warning('⏭️ 你选择了跳过此题')
                else:
                    st.error(f'**你的答案：** {ua if ua else "（未作答）"}')

                st.success(f'**正确答案：** {item["answer"]}')

                if st.button(
                    '✅ 已掌握，移出错题本',
                    key=f'rm_{item["num"]}',
                ):
                    st.session_state.wrong_book = [
                        x for x in st.session_state.wrong_book
                        if x['num'] != item['num']
                    ]
                    st.rerun()


# ╔══════════════════════════════════════════════════════════════════════════════╗
# ║  主程序入口                                                                 ║
# ╚══════════════════════════════════════════════════════════════════════════════╝

def main() -> None:
    st.set_page_config(
        page_title='镍电解工智能备考系统',
        page_icon='🔬',
        layout='wide',
        initial_sidebar_state='expanded',
    )

    # 少量全局样式
    st.markdown(
        """
        <style>
        /* 进度条圆角 */
        .stProgress > div > div { border-radius: 4px; }
        /* 标题行上边距收紧 */
        h1 { margin-top: 0.2rem !important; }
        /* expander 内容左侧对齐 */
        .streamlit-expanderContent { padding-left: 1rem; }
        </style>
        """,
        unsafe_allow_html=True,
    )

    _init_session()

    page = _render_sidebar()

    if page == _PAGE_HOME:
        _render_home()
    elif page == _PAGE_PRACTICE:
        _render_practice()
    elif page == _PAGE_WRONG:
        _render_wrong_book()


if __name__ == '__main__':
    main()
