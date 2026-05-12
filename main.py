#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
镍电解工智能备考系统 — Flet 移动端版
======================================
本地运行:        python main.py
Android APK:     flet build apk --project "镍电解工备考" --org com.nickel.exam
iOS IPA:         flet build ipa   (需 macOS + Xcode)
Windows EXE:     flet build windows

依赖安装:
    pip install flet python-docx
"""

import io
import os
import re
import difflib
import random
from collections import defaultdict
from datetime import datetime
from typing import Optional, List, Dict, Tuple

import flet as ft

try:
    from docx import Document
    DOCX_OK = True
except ImportError:
    DOCX_OK = False


# ╔══════════════════════════════════════════════════════════════════════╗
# ║  业务逻辑层（与 app.py 完全相同，零修改）                           ║
# ╚══════════════════════════════════════════════════════════════════════╝

_Q_START  = re.compile(r'^(\d{1,4})\s*[、.．。\s]\s*(.*)', re.DOTALL)
_ANS_LINE = re.compile(r'^答案?\s*[：:]\s*(.*)', re.DOTALL)


def _iter_paragraphs(doc_bytes: bytes) -> List[str]:
    """按顺序提取 docx 所有非空段落（含表格内容）。"""
    doc = Document(io.BytesIO(doc_bytes))
    lines: List[str] = []
    from docx.oxml.ns import qn

    for child in doc.element.body.iterchildren():
        tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag
        if tag == 'p':
            try:
                from docx.text.paragraph import Paragraph
                text = Paragraph(child, doc).text.strip()
            except Exception:
                text = ''
            if text:
                lines.append(text)
        elif tag == 'tbl':
            try:
                from docx.table import Table
                for row in Table(child, doc).rows:
                    for cell in row.cells:
                        for para in cell.paragraphs:
                            t = para.text.strip()
                            if t:
                                lines.append(t)
            except Exception:
                pass

    return lines or [p.text.strip() for p in doc.paragraphs if p.text.strip()]


def _check_invalid(text: str, answer: Optional[str]) -> Optional[str]:
    clean = re.sub(r'[^\w一-鿿]', '', text or '')
    if len(clean) < 3:
        return f'题干过短（净长 {len(clean)} 字）'
    if answer is None:
        return '缺少答案'
    stripped = answer.strip()
    if not stripped:
        return '答案为空'
    if re.fullmatch(r'(略|无|暂略|见课本|待定|参考答案|[\(（\)）\s/\\]+)', stripped):
        return f'答案无效（值={stripped!r}）'
    return None


def infer_type(text: str, answer: str) -> str:
    ans = answer.strip()
    if re.fullmatch(r'(对|错|√|×|T|F|true|false|正确|错误|是|否)', ans, re.IGNORECASE):
        return '判断题'
    has_opts = bool(
        re.search(r'(?m)(?:^|\s)[A-D]\s*[、.．\s。]', text) or
        re.search(r'[（(][A-D][)）]', text)
    )
    ans_letters = re.fullmatch(r'[A-Da-d]{1,4}', ans)
    if has_opts or ans_letters:
        if '多选' in text or len(re.findall(r'[A-Da-d]', ans)) > 1:
            return '多选题'
        return '选择题'
    return '填空题'


def infer_category(text: str) -> str:
    if re.search(r'安全|危险|着火|急救|液碱|酸雾|侧上风|高处作业|灭火|报警|漏电|盲板', text):
        return '安全与环保类'
    if re.search(r'疙瘩|弯板|气孔|钝化|烧板|短路|下饺子|阳极泥|打火|平板作业|掏槽', text):
        return '现场异常排查类'
    if re.search(r'\d+%|\d+℃|\d+mm|\d+V|\d+A|\d+天|密度|单耗|大于|小于', text):
        return '工艺参数硬指标'
    return '基础理论与常识'


def parse_docx(doc_bytes: bytes) -> Tuple[List[Dict], List[Dict]]:
    lines = _iter_paragraphs(doc_bytes)
    raw: List[Dict] = []
    cur_num, cur_lines, cur_ans, in_ans = None, [], None, False

    for line in lines:
        q_m = _Q_START.match(line)
        a_m = _ANS_LINE.match(line)
        if q_m:
            if cur_num is not None:
                raw.append({'num': cur_num, 'text': '\n'.join(cur_lines).strip(), 'answer': cur_ans})
            cur_num = int(q_m.group(1))
            first   = q_m.group(2).strip()
            cur_lines = [first] if first else []
            cur_ans, in_ans = None, False
        elif a_m and cur_num is not None:
            in_ans, cur_ans = True, a_m.group(1).strip()
        elif cur_num is not None:
            if in_ans:
                cur_ans = (cur_ans or '') + '\n' + line
            else:
                cur_lines.append(line)

    if cur_num is not None:
        raw.append({'num': cur_num, 'text': '\n'.join(cur_lines).strip(), 'answer': cur_ans})

    questions, error_log = [], []
    for b in raw:
        reason = _check_invalid(b['text'], b['answer'])
        if reason:
            error_log.append({'num': b['num'], 'text': (b['text'] or '')[:60], 'reason': reason})
            continue
        ans = (b['answer'] or '').strip()
        questions.append({
            'num':      b['num'],
            'text':     b['text'],
            'answer':   ans,
            'type':     infer_type(b['text'], ans),
            'category': infer_category(b['text']),
        })
    return questions, error_log


# ── 智能判分 ──────────────────────────────────────────────────────────

_BOOL_MAP = {
    '对': True, '√': True, 'T': True, 't': True, 'true': True, '正确': True, '是': True,
    '错': False,'×': False,'F': False,'f': False,'false': False,'错误': False,'否': False,
}
_PUNCT = re.compile(
    r'[\s　，、．。！？：；《》【】''"",'
    r'.;!?\'"()\[\]{}、。，；：！？《》【】]'
)


def _clean(s: str) -> str:
    return _PUNCT.sub('', s)


def smart_judge(user: str, correct: str, q_type: str) -> Tuple[bool, str]:
    user = (user or '').strip()
    if not user:
        return False, '⚠️ 未作答'
    if q_type == '判断题':
        u, c = _BOOL_MAP.get(user), _BOOL_MAP.get(correct.strip())
        if u is None:
            return False, f'⚠️ 无法识别 "{user}"，请选对/错'
        return (True, '✅ 回答正确！') if u == c else (False, '❌ 回答错误')
    if q_type in ('选择题', '多选题'):
        u = ''.join(sorted(set(re.findall(r'[A-Da-d]', user)))).upper()
        c = ''.join(sorted(set(re.findall(r'[A-Da-d]', correct)))).upper()
        if not u:
            return False, f'⚠️ 无法识别 "{user}"'
        return (True, '✅ 回答正确！') if u == c else (False, '❌ 回答错误')
    # 填空题
    u, c = _clean(user), _clean(correct)
    if u == c:
        return True, '✅ 回答正确！'
    if u.lower() == c.lower():
        return True, '✅ 回答正确！（忽略大小写）'
    ratio = difflib.SequenceMatcher(None, u, c).ratio()
    if ratio >= 0.75:
        return True, f'✅ 模糊匹配通过（相似度 {ratio:.0%}）'
    if ratio >= 0.50:
        return False, f'❌ 接近但不准确（相似度 {ratio:.0%}）'
    return False, '❌ 回答错误'


# ╔══════════════════════════════════════════════════════════════════════╗
# ║  应用状态                                                           ║
# ╚══════════════════════════════════════════════════════════════════════╝

class AppState:
    def __init__(self):
        self.questions:   List[Dict] = []
        self.error_log:   List[Dict] = []
        self.wrong_book:  List[Dict] = []
        self.loaded:      bool = False
        self.cur_q:       Optional[Dict] = None
        self.submitted:   bool = False
        self.show_ans:    bool = False
        self.last_result: Optional[Dict] = None
        self.sel_cat:     str = '全部'
        self.sel_type:    str = '全部'

    def pool(self) -> List[Dict]:
        return [
            q for q in self.questions
            if (self.sel_cat  == '全部' or q['category'] == self.sel_cat)
            and (self.sel_type == '全部' or q['type']     == self.sel_type)
        ]

    def pick(self) -> Optional[Dict]:
        p = self.pool()
        if not p:
            return None
        last = (self.cur_q or {}).get('num', -1)
        cands = [q for q in p if q['num'] != last] or p
        return random.choice(cands)

    def add_wrong(self, q: Dict, ua: str):
        if not any(x['num'] == q['num'] for x in self.wrong_book):
            self.wrong_book.append({**q, 'user_answer': ua})

    def new_question(self):
        self.cur_q      = self.pick()
        self.submitted  = False
        self.show_ans   = False
        self.last_result = None


# ╔══════════════════════════════════════════════════════════════════════╗
# ║  Flet UI                                                            ║
# ╚══════════════════════════════════════════════════════════════════════╝

# 颜色常量
C_PRIMARY  = ft.Colors.BLUE_700
C_SUCCESS  = ft.Colors.GREEN_600
C_ERROR    = ft.Colors.RED_600
C_WARNING  = ft.Colors.ORANGE_600
C_SURFACE  = ft.Colors.WHITE

TYPE_BG = {
    '填空题': ft.Colors.GREEN_700,
    '选择题': ft.Colors.BLUE_700,
    '多选题': ft.Colors.ORANGE_700,
    '判断题': ft.Colors.PURPLE_700,
}
CAT_ICON = {
    '安全与环保类':   ft.Icons.SECURITY,
    '现场异常排查类': ft.Icons.BUILD,
    '工艺参数硬指标': ft.Icons.SPEED,
    '基础理论与常识': ft.Icons.SCHOOL,
}
CAT_COLOR = {
    '安全与环保类':   ft.Colors.RED_400,
    '现场异常排查类': ft.Colors.ORANGE_400,
    '工艺参数硬指标': ft.Colors.BLUE_400,
    '基础理论与常识': ft.Colors.GREEN_400,
}


def _chip(label: str, bgcolor, fgcolor=ft.Colors.WHITE) -> ft.Container:
    """小徽章组件。"""
    return ft.Container(
        ft.Text(label, size=11, color=fgcolor, weight=ft.FontWeight.BOLD),
        bgcolor=bgcolor, border_radius=4,
        padding=ft.padding.symmetric(horizontal=8, vertical=3),
    )


def _card(content: ft.Control, margin_h=8, padding=16) -> ft.Card:
    return ft.Card(
        ft.Container(content, padding=padding),
        margin=ft.margin.symmetric(horizontal=margin_h, vertical=4),
        elevation=2,
    )


def _snack(page: ft.Page, msg: str, color=None):
    page.snack_bar = ft.SnackBar(
        ft.Text(msg, color=ft.Colors.WHITE),
        open=True,
        bgcolor=color or C_PRIMARY,
    )
    page.update()


# ── 首页大盘 ─────────────────────────────────────────────────────────

def build_home(state: AppState, on_load) -> ft.Control:
    if not state.loaded:
        return ft.Column([
            ft.Container(height=50),
            ft.Icon(ft.Icons.MENU_BOOK_ROUNDED, size=90, color=ft.Colors.BLUE_200),
            ft.Text('镍电解工智能备考系统',
                    size=20, weight=ft.FontWeight.BOLD,
                    text_align=ft.TextAlign.CENTER),
            ft.Text('请加载题库 .docx 文件',
                    size=14, color=ft.Colors.GREY_600,
                    text_align=ft.TextAlign.CENTER),
            ft.Container(height=24),
            ft.ElevatedButton(
                '  📁 选择题库文件  ',
                icon=ft.Icons.UPLOAD_FILE,
                on_click=on_load,
                style=ft.ButtonStyle(
                    bgcolor=C_PRIMARY, color=ft.Colors.WHITE,
                    padding=ft.padding.symmetric(horizontal=28, vertical=14),
                    shape=ft.RoundedRectangleBorder(radius=10),
                ),
            ),
            ft.Container(height=24),
            _card(
                ft.Column([
                    ft.Text('📖 文档格式示例', weight=ft.FontWeight.BOLD, size=14),
                    ft.Divider(height=8),
                    ft.Text(
                        '1、题目内容……\n答：答案内容\n\n'
                        '2、单选题题目（单选）\nA、选项一  B、选项二\n答：A\n\n'
                        '3、该说法是否正确？\n答：对',
                        size=12, color=ft.Colors.GREY_700,
                        font_family='Courier New',
                    ),
                ], spacing=4),
            ),
        ], horizontal_alignment=ft.CrossAxisAlignment.CENTER,
           scroll=ft.ScrollMode.AUTO, expand=True, spacing=8)

    qs, errs, wb = state.questions, state.error_log, state.wrong_book
    type_cnt: Dict[str, int] = defaultdict(int)
    cat_cnt:  Dict[str, int] = defaultdict(int)
    for q in qs:
        type_cnt[q['type']] += 1
        cat_cnt[q['category']] += 1

    def metric_tile(icon, label, value, color) -> ft.Container:
        return ft.Container(
            ft.Column([
                ft.Icon(icon, color=color, size=30),
                ft.Text(str(value), size=26, weight=ft.FontWeight.BOLD, color=color),
                ft.Text(label, size=11, color=ft.Colors.GREY_600,
                        text_align=ft.TextAlign.CENTER),
            ], horizontal_alignment=ft.CrossAxisAlignment.CENTER, spacing=4, tight=True),
            bgcolor=C_SURFACE, border_radius=12, padding=12, expand=True,
            border=ft.border.all(1, ft.Colors.GREY_200),
        )

    def prog_row(label, cnt, total, color) -> ft.Column:
        pct = cnt / total if total else 0
        return ft.Column([
            ft.Row([
                ft.Text(label, size=13, expand=True),
                ft.Text(f'{cnt} 题  {pct:.0%}', size=12, color=ft.Colors.GREY_500),
            ]),
            ft.ProgressBar(value=pct, color=color, bgcolor=ft.Colors.GREY_100, height=7,
                           border_radius=4),
            ft.Container(height=2),
        ], spacing=4)

    return ft.Column([
        ft.Container(
            ft.ElevatedButton(
                '重新加载题库', icon=ft.Icons.REFRESH,
                on_click=on_load,
                style=ft.ButtonStyle(
                    bgcolor=ft.Colors.BLUE_50, color=C_PRIMARY,
                ),
            ),
            padding=ft.padding.only(left=12, right=12, top=8),
        ),
        # 指标卡
        ft.Container(
            ft.Row([
                metric_tile(ft.Icons.MENU_BOOK, '有效题目', len(qs), C_PRIMARY),
                ft.Container(width=8),
                metric_tile(ft.Icons.BLOCK, '拦截残缺', len(errs), C_ERROR),
                ft.Container(width=8),
                metric_tile(ft.Icons.BOOKMARK, '错题本', len(wb), C_WARNING),
            ]),
            padding=ft.padding.symmetric(horizontal=12, vertical=4),
        ),
        # 题型分布
        _card(ft.Column([
            ft.Text('📊 题型分布', weight=ft.FontWeight.BOLD, size=15),
            ft.Divider(height=6),
            *[prog_row(t, c, len(qs), TYPE_BG.get(t, ft.Colors.BLUE))
              for t, c in sorted(type_cnt.items(), key=lambda x: -x[1])],
        ], spacing=2)),
        # 类别分布
        _card(ft.Column([
            ft.Text('🏷️ 类别分布', weight=ft.FontWeight.BOLD, size=15),
            ft.Divider(height=6),
            *[prog_row(cat, c, len(qs), CAT_COLOR.get(cat, ft.Colors.TEAL))
              for cat, c in sorted(cat_cnt.items(), key=lambda x: -x[1])],
        ], spacing=2)),
        # 残缺题折叠
        *([] if not errs else [
            _card(ft.Column([
                ft.Text(f'⚠️ 拦截残缺题（{len(errs)} 条）',
                        weight=ft.FontWeight.BOLD, size=14, color=C_WARNING),
                ft.Divider(height=6),
                *[ft.Text(f'• 第{e["num"]}题  {e["reason"]}  {e["text"]}',
                          size=12, color=ft.Colors.GREY_600)
                  for e in errs[:30]],
                *([] if len(errs) <= 30 else
                  [ft.Text(f'…还有 {len(errs)-30} 条', size=11, color=ft.Colors.GREY_500)]),
            ], spacing=4)),
        ]),
    ], scroll=ft.ScrollMode.AUTO, expand=True, spacing=0)


# ── 智能刷题 ─────────────────────────────────────────────────────────

def build_practice(state: AppState, refresh) -> ft.Control:
    if not state.loaded:
        return ft.Column([
            ft.Container(height=80),
            ft.Icon(ft.Icons.QUIZ_OUTLINED, size=70, color=ft.Colors.BLUE_100),
            ft.Text('请先在首页加载题库', color=ft.Colors.GREY_500,
                    text_align=ft.TextAlign.CENTER),
        ], horizontal_alignment=ft.CrossAxisAlignment.CENTER, expand=True)

    qs = state.questions
    all_cats  = ['全部'] + sorted({q['category'] for q in qs})
    all_types = ['全部'] + sorted({q['type']     for q in qs})

    # 答案输入持有器（跨回调共享）
    ans_holder = {'val': ''}

    def on_cat_change(e):
        state.sel_cat = e.control.value
        state.new_question()
        refresh()

    def on_type_change(e):
        state.sel_type = e.control.value
        state.new_question()
        refresh()

    def on_pick(_):
        state.new_question()
        refresh()

    def on_submit(_):
        ua = ans_holder['val'].strip()
        if not ua:
            return
        ok, msg = smart_judge(ua, state.cur_q['answer'], state.cur_q['type'])
        state.last_result = {'ok': ok, 'msg': msg, 'user_ans': ua}
        state.submitted = True
        if not ok:
            state.add_wrong(state.cur_q, ua)
        refresh()

    def on_skip(_):
        state.show_ans = True
        state.add_wrong(state.cur_q, '（跳过）')
        refresh()

    def on_next(_):
        state.new_question()
        refresh()

    pool = state.pool()

    # 筛选行
    filter_row = ft.Row([
        ft.Dropdown(
            label='类别', value=state.sel_cat, dense=True, expand=True,
            options=[ft.dropdown.Option(c) for c in all_cats],
            on_change=on_cat_change,
        ),
        ft.Dropdown(
            label='题型', value=state.sel_type, dense=True, expand=True,
            options=[ft.dropdown.Option(t) for t in all_types],
            on_change=on_type_change,
        ),
    ], spacing=8)

    tool_row = ft.Row([
        ft.Text(f'范围 {len(pool)} 题', size=12, color=ft.Colors.GREY_500, expand=True),
        ft.ElevatedButton(
            '🎲 随机抽题', on_click=on_pick,
            style=ft.ButtonStyle(bgcolor=C_PRIMARY, color=ft.Colors.WHITE),
        ),
    ])

    if not pool:
        return ft.Column([
            ft.Container(ft.Column([filter_row, tool_row], spacing=8),
                         padding=ft.padding.symmetric(horizontal=12, vertical=8)),
            ft.Container(
                ft.Text('当前筛选下无题目，请调整', color=ft.Colors.GREY_500,
                        text_align=ft.TextAlign.CENTER),
                alignment=ft.alignment.center, expand=True,
            ),
        ], expand=True)

    if state.cur_q is None:
        state.cur_q = state.pick()

    q = state.cur_q

    # 输入控件
    if q['type'] == '判断题':
        ans_holder['val'] = '对'  # 默认值

        def on_radio(e):
            ans_holder['val'] = '对' if e.control.value == 'T' else '错'

        input_ctrl = ft.RadioGroup(
            value='T',
            content=ft.Row([
                ft.Radio(value='T', label='对  ✓'),
                ft.Container(width=20),
                ft.Radio(value='F', label='错  ✗'),
            ]),
            on_change=on_radio,
        )

    elif q['type'] in ('选择题', '多选题'):
        hint = '输入字母（如 A 或 ABC，顺序不限）'

        def on_tf_change(e):
            ans_holder['val'] = e.control.value

        input_ctrl = ft.TextField(
            label=hint, on_change=on_tf_change, autofocus=False,
            text_style=ft.TextStyle(size=20, weight=ft.FontWeight.BOLD),
            text_align=ft.TextAlign.CENTER,
            capitalization=ft.TextCapitalization.CHARACTERS,
            border_radius=8,
        )
    else:
        def on_fb_change(e):
            ans_holder['val'] = e.control.value

        input_ctrl = ft.TextField(
            label='请输入答案', multiline=True, min_lines=2, max_lines=4,
            on_change=on_fb_change, border_radius=8,
        )

    # 反馈区
    feedback: List[ft.Control] = []

    if state.submitted and state.last_result:
        res = state.last_result
        ok  = res['ok']
        feedback = [
            ft.Container(
                ft.Row([
                    ft.Icon(ft.Icons.CHECK_CIRCLE if ok else ft.Icons.CANCEL,
                            color=C_SUCCESS if ok else C_ERROR),
                    ft.Text(res['msg'], color=C_SUCCESS if ok else C_ERROR,
                            size=14, weight=ft.FontWeight.BOLD),
                ], spacing=8),
                bgcolor=ft.Colors.GREEN_50 if ok else ft.Colors.RED_50,
                border=ft.border.all(1, C_SUCCESS if ok else C_ERROR),
                border_radius=8, padding=10,
            ),
        ]
        if not ok:
            feedback.append(ft.Container(
                ft.Row([
                    ft.Icon(ft.Icons.LIGHTBULB, color=C_WARNING),
                    ft.Text(f'标准答案：{q["answer"]}', size=14,
                            weight=ft.FontWeight.BOLD, expand=True),
                ], spacing=8),
                bgcolor=ft.Colors.AMBER_50, border_radius=8, padding=10,
            ))
        feedback.append(ft.ElevatedButton(
            '➡️  下一题', on_click=on_next, expand=True,
            style=ft.ButtonStyle(
                bgcolor=C_PRIMARY, color=ft.Colors.WHITE,
                padding=ft.padding.symmetric(vertical=12),
                shape=ft.RoundedRectangleBorder(radius=8),
            ),
        ))

    elif state.show_ans:
        feedback = [
            ft.Container(
                ft.Row([
                    ft.Icon(ft.Icons.LIGHTBULB, color=C_WARNING),
                    ft.Text(f'标准答案：{q["answer"]}', size=14,
                            weight=ft.FontWeight.BOLD, expand=True),
                ], spacing=8),
                bgcolor=ft.Colors.AMBER_50, border_radius=8, padding=10,
            ),
            ft.ElevatedButton(
                '➡️  下一题', on_click=on_next, expand=True,
                style=ft.ButtonStyle(
                    bgcolor=C_PRIMARY, color=ft.Colors.WHITE,
                    padding=ft.padding.symmetric(vertical=12),
                    shape=ft.RoundedRectangleBorder(radius=8),
                ),
            ),
        ]
    else:
        feedback = [
            ft.Row([
                ft.ElevatedButton(
                    '✅ 提交', on_click=on_submit, expand=True,
                    style=ft.ButtonStyle(
                        bgcolor=C_SUCCESS, color=ft.Colors.WHITE,
                        padding=ft.padding.symmetric(vertical=12),
                        shape=ft.RoundedRectangleBorder(radius=8),
                    ),
                ),
                ft.OutlinedButton(
                    '👁 看答案', on_click=on_skip, expand=True,
                    style=ft.ButtonStyle(
                        padding=ft.padding.symmetric(vertical=12),
                        shape=ft.RoundedRectangleBorder(radius=8),
                    ),
                ),
            ], spacing=8),
        ]

    # 题目卡片
    q_card = _card(ft.Column([
        ft.Row([
            _chip(q['type'], TYPE_BG.get(q['type'], C_PRIMARY)),
            _chip(q['category'], CAT_COLOR.get(q['category'], ft.Colors.TEAL)),
            ft.Text(f'# {q["num"]}', size=11, color=ft.Colors.GREY_400),
        ], spacing=6, wrap=True),
        ft.Divider(height=10),
        ft.Text(q['text'], size=16, selectable=True),
        ft.Divider(height=10),
        input_ctrl,
        ft.Container(height=6),
        *feedback,
    ], spacing=8))

    return ft.Column([
        ft.Container(
            ft.Column([filter_row, tool_row], spacing=8),
            padding=ft.padding.symmetric(horizontal=12, vertical=8),
        ),
        q_card,
    ], scroll=ft.ScrollMode.AUTO, expand=True, spacing=0)


# ── 错题本 ────────────────────────────────────────────────────────────

def build_wrong_book(state: AppState, refresh, page: ft.Page) -> ft.Control:
    wb = state.wrong_book

    if not wb:
        return ft.Column([
            ft.Container(height=80),
            ft.Icon(ft.Icons.TASK_ALT, size=70, color=ft.Colors.GREEN_300),
            ft.Text('错题本为空，继续加油！',
                    color=ft.Colors.GREY_500, text_align=ft.TextAlign.CENTER),
        ], horizontal_alignment=ft.CrossAxisAlignment.CENTER, expand=True)

    def on_export(_):
        lines = [
            '镍电解工智能备考系统 · 错题本导出',
            f'共 {len(wb)} 道错题  导出时间：{datetime.now().strftime("%Y-%m-%d %H:%M")}',
            '=' * 60, '',
        ]
        for i, q in enumerate(wb, 1):
            lines += [
                f'第 {i} 题（原题号 {q["num"]}）',
                f'题型：{q["type"]}  |  类别：{q["category"]}',
                f'题目：{q["text"]}',
                f'你的答案：{q.get("user_answer", "")}',
                f'正确答案：{q["answer"]}',
                '-' * 60, '',
            ]
        content = '\n'.join(lines)

        # 保存路径：优先 Downloads，其次当前目录
        candidates = [
            os.path.expanduser('~/Downloads'),
            os.path.expanduser('~/Documents'),
            '/sdcard/Download',
            os.getcwd(),
        ]
        save_dir = next((d for d in candidates if os.path.isdir(d)), os.getcwd())
        ts       = datetime.now().strftime('%Y%m%d_%H%M%S')
        path     = os.path.join(save_dir, f'错题本_{ts}.txt')
        try:
            with open(path, 'w', encoding='utf-8-sig') as f:
                f.write(content)
            _snack(page, f'✅ 已保存至：{path}', ft.Colors.GREEN_700)
        except Exception as exc:
            _snack(page, f'❌ 保存失败：{exc}', C_ERROR)

    def on_clear(_):
        def confirm(_):
            dlg.open = False
            state.wrong_book.clear()
            page.update()
            refresh()

        def cancel(_):
            dlg.open = False
            page.update()

        dlg = ft.AlertDialog(
            modal=True,
            title=ft.Text('确认清空错题本？'),
            content=ft.Text('清空后无法恢复'),
            actions=[
                ft.TextButton('取消', on_click=cancel),
                ft.ElevatedButton(
                    '确认清空',
                    on_click=confirm,
                    style=ft.ButtonStyle(bgcolor=C_ERROR, color=ft.Colors.WHITE),
                ),
            ],
        )
        page.overlay.append(dlg)
        dlg.open = True
        page.update()

    def make_item(item: Dict) -> ft.Card:
        ua      = item.get('user_answer', '')
        skipped = (ua == '（跳过）')

        def on_remove(_):
            state.wrong_book[:] = [x for x in state.wrong_book if x['num'] != item['num']]
            refresh()

        return ft.Card(
            ft.Container(
                ft.Column([
                    ft.Row([
                        _chip(item['type'], TYPE_BG.get(item['type'], C_PRIMARY)),
                        ft.Text(f'# {item["num"]}', size=11, color=ft.Colors.GREY_400),
                    ], spacing=6),
                    ft.Text(item['text'], size=14, selectable=True),
                    ft.Divider(height=4),
                    ft.Container(
                        ft.Text(
                            '⏭ 已跳过' if skipped else f'你的答案：{ua}',
                            size=13,
                            color=C_WARNING if skipped else C_ERROR,
                        ),
                        bgcolor=ft.Colors.ORANGE_50 if skipped else ft.Colors.RED_50,
                        border_radius=6, padding=8,
                    ),
                    ft.Container(
                        ft.Text(f'正确答案：{item["answer"]}', size=13,
                                color=C_SUCCESS, weight=ft.FontWeight.BOLD),
                        bgcolor=ft.Colors.GREEN_50, border_radius=6, padding=8,
                    ),
                    ft.TextButton(
                        '✅ 已掌握，移出错题本',
                        on_click=on_remove,
                        style=ft.ButtonStyle(color=C_SUCCESS),
                    ),
                ], spacing=6),
                padding=14,
            ),
            margin=ft.margin.symmetric(horizontal=8, vertical=3),
            elevation=1,
        )

    grouped: Dict[str, List[Dict]] = defaultdict(list)
    for item in wb:
        grouped[item['category']].append(item)

    items_ctrl: List[ft.Control] = []
    for cat, items in grouped.items():
        items_ctrl.append(ft.Container(
            ft.Row([
                ft.Icon(CAT_ICON.get(cat, ft.Icons.LABEL),
                        size=16, color=CAT_COLOR.get(cat, ft.Colors.TEAL)),
                ft.Text(f'{cat}（{len(items)} 题）',
                        weight=ft.FontWeight.BOLD, size=14),
            ], spacing=6),
            padding=ft.padding.only(left=14, top=10, bottom=2),
        ))
        for item in items:
            items_ctrl.append(make_item(item))

    return ft.Column([
        ft.Container(
            ft.Row([
                ft.Text(f'共 {len(wb)} 道错题',
                        weight=ft.FontWeight.BOLD, size=15, expand=True),
                ft.OutlinedButton(
                    '导出', icon=ft.Icons.DOWNLOAD, on_click=on_export,
                    style=ft.ButtonStyle(color=C_PRIMARY),
                ),
                ft.Container(width=4),
                ft.OutlinedButton(
                    '清空', icon=ft.Icons.DELETE_OUTLINE, on_click=on_clear,
                    style=ft.ButtonStyle(color=C_ERROR),
                ),
            ]),
            padding=ft.padding.symmetric(horizontal=12, vertical=8),
        ),
        ft.Divider(height=1),
        *items_ctrl,
    ], scroll=ft.ScrollMode.AUTO, expand=True, spacing=0)


# ╔══════════════════════════════════════════════════════════════════════╗
# ║  主函数入口                                                         ║
# ╚══════════════════════════════════════════════════════════════════════╝

def main(page: ft.Page):
    page.title        = '镍电解工备考系统'
    page.theme_mode   = ft.ThemeMode.LIGHT
    page.theme        = ft.Theme(color_scheme_seed=ft.Colors.BLUE)
    page.padding      = 0
    page.bgcolor      = ft.Colors.GREY_50
    page.window.width  = 420   # 模拟手机宽度（桌面调试用）
    page.window.height = 820

    state = AppState()

    # ── 内容区 ────────────────────────────────────────────────────
    content = ft.Column([], expand=True)

    def refresh():
        idx = nav.selected_index
        if idx == 0:
            content.controls = [build_home(state, on_load)]
        elif idx == 1:
            content.controls = [build_practice(state, refresh)]
        elif idx == 2:
            content.controls = [build_wrong_book(state, refresh, page)]
        page.update()

    # ── 文件选择器 ────────────────────────────────────────────────
    def on_file_result(e: ft.FilePickerResultEvent):
        if not e.files:
            return
        path = e.files[0].path
        if not path:
            _snack(page, '无法获取文件路径', C_ERROR)
            return

        # 加载弹窗
        dlg = ft.AlertDialog(
            modal=True,
            content=ft.Column([
                ft.ProgressRing(width=40, height=40),
                ft.Text('正在解析题库…', text_align=ft.TextAlign.CENTER),
            ], horizontal_alignment=ft.CrossAxisAlignment.CENTER,
               tight=True, spacing=12),
        )
        page.overlay.append(dlg)
        dlg.open = True
        page.update()

        try:
            with open(path, 'rb') as f:
                raw = f.read()
            qs, errs = parse_docx(raw)
            state.questions  = qs
            state.error_log  = errs
            state.wrong_book = []
            state.loaded     = True
            state.cur_q      = None
            state.submitted  = False
            state.show_ans   = False
            state.last_result = None
            dlg.open = False
            _snack(page,
                   f'✅ 解析完成：{len(qs)} 道有效题，拦截 {len(errs)} 道残缺题',
                   ft.Colors.GREEN_700)
        except Exception as exc:
            dlg.open = False
            _snack(page, f'❌ 解析失败：{exc}', C_ERROR)
        finally:
            page.update()
            refresh()

    picker = ft.FilePicker(on_result=on_file_result)
    page.overlay.append(picker)

    def on_load(_=None):
        picker.pick_files(
            allowed_extensions=['docx'],
            dialog_title='选择题库文件',
        )

    # ── 底部导航 ──────────────────────────────────────────────────
    def on_nav(e):
        refresh()

    nav = ft.NavigationBar(
        destinations=[
            ft.NavigationBarDestination(
                icon=ft.Icons.HOME_OUTLINED,
                selected_icon=ft.Icons.HOME,
                label='首页',
            ),
            ft.NavigationBarDestination(
                icon=ft.Icons.QUIZ_OUTLINED,
                selected_icon=ft.Icons.QUIZ,
                label='刷题',
            ),
            ft.NavigationBarDestination(
                icon=ft.Icons.BOOKMARK_OUTLINE,
                selected_icon=ft.Icons.BOOKMARK,
                label='错题本',
            ),
        ],
        on_change=on_nav,
        selected_index=0,
        bgcolor=ft.Colors.WHITE,
        indicator_color=ft.Colors.BLUE_50,
    )

    page.appbar = ft.AppBar(
        title=ft.Text('🔬 镍电解工备考系统',
                      color=ft.Colors.WHITE, size=16, weight=ft.FontWeight.BOLD),
        bgcolor=C_PRIMARY,
        center_title=False,
        actions=[
            ft.IconButton(
                ft.Icons.UPLOAD_FILE,
                icon_color=ft.Colors.WHITE,
                tooltip='加载题库',
                on_click=on_load,
            ),
        ],
    )
    page.navigation_bar = nav
    page.add(content)

    refresh()   # 渲染首页


if __name__ == '__main__':
    ft.app(target=main)
